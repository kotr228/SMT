#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "Times New Roman"
SIZE = 14

def set_font(run, bold=False, italic=False, size=SIZE):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)
    existing = rPr.find(qn('w:rFonts'))
    if existing is not None:
        rPr.remove(existing)
    rPr.insert(0, rFonts)

def add_paragraph(doc, text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=SIZE):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run(text)
    set_font(run, bold=bold, italic=italic, size=size)
    return p

def add_bullet(doc, text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Cm(1.0)
    run = p.add_run("— " + text)
    set_font(run, bold=bold)
    return p

def add_empty(doc):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run("")
    set_font(run)
    return p

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

# Title
add_paragraph(doc, "Семінарське заняття №4", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_empty(doc)

# Section 1
add_paragraph(doc, "1. Технічна документація ІТ-проекту (тема КП):", bold=True)
add_empty(doc)

add_paragraph(doc, "Назва проекту (тема КП): DocControlSolution — розподілена система контролю версій та управління документами для корпоративної мережі.")
add_empty(doc)

add_paragraph(doc, "Технічні специфікації:", bold=True)
specs = [
    "мова програмування: C# (.NET 8, платформа net8.0-windows);",
    "технологія бекенду: Windows Service (System.ServiceProcess.ServiceController);",
    "технологія графічного інтерфейсу (GUI): WPF (Windows Presentation Foundation) з використанням XAML-розмітки та бібліотеки MahApps.Metro 2.4.11;",
    "міжпроцесна комунікація (IPC) між сервісом та UI: Named Pipes (System.IO.Pipes);",
    "бібліотека контролю версій: LibGit2Sharp 0.31.0 (обгортка над libgit2 для роботи з Git-репозиторіями);",
    "база даних: SQLite через Microsoft.Data.Sqlite 9.0.9;",
    "виявлення вузлів у мережі: UDP Broadcast (System.Net.Sockets.UdpClient);",
    "передача файлів між вузлами: TCP (власний бінарний протокол);",
    "AI-аналіз структури директорій: Ollama (локальна мовна модель, HTTP API);",
    "картографічний модуль: Google Maps, Bing Maps, OpenStreetMap (WebView2 + map.html);",
    "допоміжні бібліотеки: ClosedXML 0.104.1 (експорт до Excel), DocumentFormat.OpenXml 3.0.2, Microsoft.Web.WebView2 1.0.3537.50.",
]
for s in specs:
    add_bullet(doc, s)

add_empty(doc)
add_paragraph(doc, "Функціональність:", bold=True)
funcs = [
    "моніторинг директорій локальної файлової системи — відстеження нових файлів, змін та видалень у заданих папках;",
    "автоматичне версіонування файлів через Git (LibGit2Sharp) із плановими комітами (за замовчуванням кожні 12 годин);",
    "мережеве виявлення вузлів у локальній мережі через UDP broadcast із збереженням списку активних пристроїв;",
    "перегляд файлів та директорій на віддалених вузлах і їх завантаження через TCP;",
    "блокування файлів для ексклюзивного редагування (file lock) з реєстрацією у базі даних;",
    "IP-фільтрація для обмеження доступу до мережевих ресурсів;",
    "AI-аналіз структури директорій та генерація хронологічних і географічних дорожніх карт (Ollama);",
    "географічне картографування дорожніх карт на інтерактивній карті (Google/Bing/OpenStreetMap);",
    "логування помилок та комітів у базі даних SQLite;",
    "підтримка зовнішніх сервісів та пристроїв (ExternalService, DeviceRepository);",
    "експорт даних до форматів Excel та інших документів.",
]
for f in funcs:
    add_bullet(doc, f)

add_empty(doc)
add_paragraph(doc, "Архітектурні діаграми:", bold=True)
add_empty(doc)
add_paragraph(doc, "Проект реалізовано у вигляді багатошарового рішення (solution) з шістьма окремими проектами, між якими чітко розподілено відповідальності:")
add_empty(doc)

arch_rows = [
    ("DocControlService", "Windows Service — основний бекенд: моніторинг файлів, SQLite-БД, версіонування, IP-фільтрація, Named Pipes-сервер."),
    ("DocControlUI", "WPF-додаток — графічний інтерфейс користувача; комунікує з сервісом через Named Pipes."),
    ("DocControlNetworkCore", "Мережевий шар — UDP-виявлення вузлів (DiscoveryService), TCP-передача файлів (FileTransferService), безпека (SecurityService)."),
    ("DocControlAI", "AI-модуль — підключення до Ollama, аналіз структури директорій, генерація roadmap (ChronologicalRoadmapGenerator, GeoRoadmapGenerator)."),
    ("DocControl.Maps.Core", "Картографічний модуль — провайдери карт (Google, Bing, OpenStreetMap), геокодування, офлайн-кеш тайлів."),
    ("DocControlService.Shared", "Спільні інтерфейси (IFileSystemService) та моделі (RemoteNode, PeerIdentity, FileSystemItem) для використання між проектами."),
]

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
# Header row
hdr = table.rows[0].cells
hdr[0].text = "Проект"
hdr[1].text = "Призначення"
for cell in hdr:
    for para in cell.paragraphs:
        for run in para.runs:
            set_font(run, bold=True)

for proj, desc in arch_rows:
    row = table.add_row().cells
    row[0].text = proj
    row[1].text = desc
    for cell in row:
        for para in cell.paragraphs:
            for run in para.runs:
                set_font(run)

add_empty(doc)
add_empty(doc)

# Structural choice
add_paragraph(doc, "Вибір структури проекту:", bold=True)
add_empty(doc)
add_paragraph(doc,
    "Для даного проекту обрано сервіс-орієнтовану архітектуру з чітким розподілом рівнів (backend-service + UI + network + AI + maps), оскільки система призначена для тривалої фонової роботи у корпоративній мережі. "
    "Windows Service забезпечує автозапуск при старті ОС та незалежну від входу користувача роботу. "
    "Комунікація через Named Pipes обрана як найефективніший IPC-механізм для Windows з мінімальною затримкою у межах однієї машини. "
    "Окремий мережевий проект (DocControlNetworkCore) дозволяє легко замінювати або розширювати протоколи (UDP/TCP) без зміни бізнес-логіки. "
    "Використання SQLite спрощує розгортання — не потребує встановлення окремого сервера БД. "
    "LibGit2Sharp обрано для версіонування, оскільки Git є стандартом у галузі та надає повноцінну історію змін без додаткових витрат."
)

add_empty(doc)

# Usage instructions
add_paragraph(doc, "Інструкція з використання:", bold=True)
add_empty(doc)
steps = [
    "запустити DocControlService.exe від імені адміністратора (програма автоматично запитує підвищення прав через UAC); за необхідності відлагодження додати аргумент --debug для запуску у консольному режимі;",
    "запустити DocControlUI.exe — застосунок автоматично підключається до сервісу через Named Pipes; у разі недоступності сервісу відобразиться повідомлення про помилку підключення;",
    "на вкладці «Директорії» додати папки для моніторингу — вказати шлях, налаштувати права доступу та мережеві дозволи;",
    "на вкладці «Пристрої» переглянути список виявлених у локальній мережі вузлів (оновлюється автоматично кожні 15 секунд) та налаштувати IP-фільтри для обмеження доступу;",
    "для перегляду файлів на віддаленому вузлі — обрати пристрій зі списку, натиснути «Переглянути файли»; для завантаження файлу — вибрати його у переліку та натиснути «Завантажити»;",
    "у вкладці «Блокування» переглянути заблоковані файли та за потреби зняти блокування;",
    "у вкладці «AI-аналіз» запустити аналіз структури директорій (потребує встановленого та запущеного Ollama-сервера); переглянути згенеровані дорожні карти на вкладці «Карта»;",
    "у вкладці «Журнал» переглянути історію автоматичних комітів та зареєстровані помилки.",
]
for i, step in enumerate(steps, 1):
    add_bullet(doc, f"{i}. {step}")

add_empty(doc)

# Strong points
add_paragraph(doc, "2. Сильні сторони проекту:", bold=True)
strengths = [
    "чітка багатошарова архітектура — розподіл на окремі проекти (Service/UI/Network/AI/Maps) спрощує підтримку та масштабування системи;",
    "автоматичне версіонування файлів через Git без участі користувача — будь-які зміни фіксуються та можуть бути відновлені з повною історією;",
    "P2P-виявлення вузлів у локальній мережі без центрального сервера — знижує точку відмови та спрощує розгортання;",
    "інтеграція AI (Ollama) для інтелектуального аналізу структури файлів та генерації рекомендацій;",
    "підтримка географічних дорожніх карт з відображенням на інтерактивних картах (Google Maps, Bing, OSM);",
    "механізм блокування файлів попереджає конфлікти при спільній роботі кількох користувачів;",
    "гнучке управління доступом через IP-фільтрацію — можливість обмежити доступ до ресурсів на рівні мережевих адрес.",
]
for s in strengths:
    add_bullet(doc, s)

add_empty(doc)

# Weak points
add_paragraph(doc, "Слабкі сторони проекту:", bold=True)
weaknesses = [
    "прив'язка до платформи Windows (net8.0-windows, WPF, Windows Service) — унеможливлює використання на Linux/macOS;",
    "обов'язкова наявність прав адміністратора для запуску сервісу — ускладнює розгортання у корпоративних середовищах із жорсткою політикою безпеки;",
    "залежність від локально встановленого Ollama-сервера для AI-функцій — потребує додаткового налаштування та значних апаратних ресурсів;",
    "UDP broadcast не перетинає маршрутизатори (router), тому виявлення вузлів обмежене одним мережевим сегментом (subnet);",
    "Named Pipes для IPC-комунікації між сервісом та UI означає, що UI-додаток може підключитися лише на тій самій машині, де запущено сервіс;",
    "відсутня автентифікація та шифрування даних у Named Pipes IPC — потенційна вразливість для привілейованих процесів.",
]
for w in weaknesses:
    add_bullet(doc, w)

add_empty(doc)

# Potential problems
add_paragraph(doc, "Потенційні проблеми у програмі:", bold=True)
problems = [
    "необмежене зростання Git-репозиторіїв з часом — у директоріях з великою кількістю великих файлів репозиторій може досягнути критичного розміру, що сповільнить роботу LibGit2Sharp та збільшить час коміту;",
    "перевантаження мережі у разі великої кількості одночасних вузлів — регулярні UDP broadcast кожні 10 секунд від багатьох пристроїв можуть створювати зайве навантаження;",
    "відсутність механізму вирішення конфліктів версій — якщо два користувачі одночасно редагують один файл (навіть з блокуванням), автоматичне злиття змін не передбачено;",
    "залежність від зовнішнього Ollama-сервера — у разі його недоступності або оновлення API всі AI-функції перестають працювати;",
    "SQLite може стати вузьким місцем при інтенсивному паралельному записі даних від багатьох вузлів мережі через однофайловий характер бази даних.",
]
for p in problems:
    add_bullet(doc, p)

# Save
out = "/home/user/SMT/Групова динаміка/Семінар 4 Групова динаміка.docx"
doc.save(out)
print(f"Saved: {out}")
