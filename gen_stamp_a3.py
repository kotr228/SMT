"""
Generate blank A3 landscape document with large engineering drawing title block (великий штамп)
per ДСТУ ГОСТ 2.104:2006 Form 1, adapted for full-width A3 document.
"""

from docx import Document
from docx.shared import Cm, Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_font(run, bold=False, size=8, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), "Times New Roman")
    rFonts.set(qn('w:hAnsi'), "Times New Roman")
    rFonts.set(qn('w:cs'), "Times New Roman")
    existing = rPr.find(qn('w:rFonts'))
    if existing is not None:
        rPr.remove(existing)
    rPr.insert(0, rFonts)


def cell_set(cell, text, bold=False, size=8,
             halign=WD_ALIGN_PARAGRAPH.CENTER,
             valign=WD_ALIGN_VERTICAL.CENTER):
    cell.vertical_alignment = valign
    # Clear existing paragraphs
    for p in cell.paragraphs:
        p.clear()
    p = cell.paragraphs[0]
    p.alignment = halign
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if text:
        run = p.add_run(text)
        set_font(run, bold=bold, size=size)


def set_tcW(cell, width_cm):
    """Set exact cell width in twips (1 cm ≈ 567 twips)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    twips = round(width_cm * 566.93)
    tcW.set(qn('w:w'), str(twips))
    tcW.set(qn('w:type'), 'dxa')
    existing = tcPr.find(qn('w:tcW'))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(tcW)


def set_tbl_width(table, width_cm):
    """Set total table width."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(round(width_cm * 566.93)))
    tblW.set(qn('w:type'), 'dxa')
    existing = tblPr.find(qn('w:tblW'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(tblW)


def set_no_margin_para(p):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(10)


def remove_cell_top_border(cell):
    """Remove top border of a cell (for seamless frame line)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'nil')
    existing = tcBorders.find(qn('w:top'))
    if existing is not None:
        tcBorders.remove(existing)
    tcBorders.append(top)


def make_stamp_a3():
    doc = Document()

    # ── A3 landscape page settings ────────────────────────────────────────────
    section = doc.sections[0]
    section.page_width  = Mm(420)
    section.page_height = Mm(297)
    section.left_margin   = Mm(20)
    section.right_margin  = Mm(5)
    section.top_margin    = Mm(5)
    section.bottom_margin = Mm(5)

    # Remove Normal style spacing
    normal = doc.styles['Normal']
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after  = Pt(0)

    # ── Drawing area (blank space) ────────────────────────────────────────────
    # Content height = 297-5-5 = 287mm; stamp height ≈ 55mm (5×11mm rows)
    # Drawing area ≈ 232mm.  At 10pt line spacing ≈ 3.53mm/line → ~65 lines
    for _ in range(60):
        p = doc.add_paragraph()
        set_no_margin_para(p)

    # ── Stamp table ───────────────────────────────────────────────────────────
    # Content width = 420-20-5 = 395mm = 39.5 cm
    #
    # Column layout (14 cols, total 39.5 cm):
    #  Change section (cols 0-4)  : 0.7+0.7+1.0+1.5+1.1 = 5.0 cm
    #  Personnel section (cols 5-8): 2.5+3.5+1.5+1.5    = 9.0 cm
    #  Title (col 9)               :                      18.0 cm
    #  Right section (cols 10-13)  : 0.7+1.0+1.5+4.3    = 7.5 cm
    #  TOTAL                                              39.5 cm
    #
    col_w = [
        0.7, 0.7, 1.0, 1.5, 1.1,   # 0-4  change
        2.5, 3.5, 1.5, 1.5,          # 5-8  personnel
        18.0,                          # 9    title
        0.7, 1.0, 1.5, 4.3,           # 10-13 right
    ]
    assert abs(sum(col_w) - 39.5) < 0.01, f"Width mismatch: {sum(col_w)}"

    N_ROWS, N_COLS = 5, 14
    table = doc.add_table(rows=N_ROWS, cols=N_COLS)
    table.style = 'Table Grid'
    set_tbl_width(table, 39.5)

    ROW_H = 1.1  # cm per row  (5 × 11mm = 55mm stamp height)

    for r in range(N_ROWS):
        row = table.rows[r]
        row.height      = Cm(ROW_H)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for c in range(N_COLS):
            set_tcW(row.cells[c], col_w[c])

    # ── Populate cells before merging ────────────────────────────────────────

    # --- Row 0: column headers (change section) + first personnel row ----------
    cell_set(table.cell(0, 0), 'Зм.',      size=7)
    cell_set(table.cell(0, 1), 'Арк.',     size=7)
    cell_set(table.cell(0, 2), '№докум.',  size=7)
    cell_set(table.cell(0, 3), 'Підпис',   size=7)
    cell_set(table.cell(0, 4), 'Дата',     size=7)
    cell_set(table.cell(0, 5), 'Розробив', size=8)
    cell_set(table.cell(0, 6), '',         size=8)   # name blank
    cell_set(table.cell(0, 7), 'Підпис',   size=7)
    cell_set(table.cell(0, 8), 'Дата',     size=7)
    # col 9 → title (will be merged rows 0-3)
    cell_set(table.cell(0, 10), 'Літ.',     size=7)
    cell_set(table.cell(0, 11), 'Аркуш',   size=7)
    cell_set(table.cell(0, 12), 'Аркушів', size=7)
    # col 13 → dept/faculty (will be merged rows 0-3)

    # --- Row 1: Перевірив -------------------------------------------------------
    cell_set(table.cell(1, 5), 'Перевірив', size=8)

    # --- Row 2: T.контр. (технічний контроль) ----------------------------------
    cell_set(table.cell(2, 5), 'Т.контр.',  size=8)

    # --- Row 3: Н.контр. -------------------------------------------------------
    cell_set(table.cell(3, 5), 'Н.контр.',  size=8)

    # --- Row 4: Затв. ----------------------------------------------------------
    cell_set(table.cell(4, 5), 'Затв.',     size=8)

    # ── Merge title block: col 9, rows 0-3 ───────────────────────────────────
    table.cell(0, 9).merge(table.cell(3, 9))
    # Set title area text (blank – user will type)
    cell_set(table.cell(0, 9), '', size=14, bold=True)

    # ── Merge dept/faculty block: col 13, rows 0-3 ───────────────────────────
    table.cell(0, 13).merge(table.cell(3, 13))
    # Leave blank for user to fill in faculty/department
    cell_set(table.cell(0, 13), '', size=8)

    # ── Row 4 right side: merge cols 9-13 for organization name ─────────────
    table.cell(4, 9).merge(table.cell(4, 13))
    cell_set(table.cell(4, 9), '', size=8, halign=WD_ALIGN_PARAGRAPH.LEFT)

    # ── Right section labels in rows 1-3 (Літ value, Аркуш №, Аркушів) ──────
    # Row 1: actual sheet number and total go in cols 11-12
    cell_set(table.cell(1, 10), '', size=8)   # Літера value
    cell_set(table.cell(1, 11), '', size=8)   # Аркуш number
    cell_set(table.cell(1, 12), '', size=8)   # Аркушів total
    # Rows 2-3 in cols 10-12: leave blank
    for r in (2, 3):
        for c in (10, 11, 12):
            cell_set(table.cell(r, c), '', size=8)

    # ── Save ──────────────────────────────────────────────────────────────────
    out_path = '/home/user/SMT/Штамп А3 альбомний.docx'
    doc.save(out_path)
    print(f'Saved: {out_path}')


if __name__ == '__main__':
    make_stamp_a3()
